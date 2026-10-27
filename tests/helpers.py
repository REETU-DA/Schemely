"""Test helpers: build datasets/analyses from in-memory text."""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from modules import workflow as wf  # noqa: E402
from modules.file_reader import read_file  # noqa: E402


def csv_bytes(text: str) -> bytes:
    return text.strip("\n").replace("\r\n", "\n").encode("utf-8")


def analyse_csvs(files: dict, dialect="postgresql", mode="direct", staging=False, accept_all=False):
    pairs = [(n, csv_bytes(t)) for n, t in files.items()]
    ds, reps = wf.load(pairs)
    assert all(r.status != "error" for r in reps), [r.errors for r in reps]
    a = wf.analyse(ds, dialect=dialect, mode=mode, staging=staging)
    if accept_all:
        for fk in a.schema.fks:
            fk.accepted = True
        wf.revalidate(a)
    return a


def xlsx_bytes(sheets: dict) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def issues(a, code):
    return [i for i in a.validation.issues if i.code == code]


def table(a, name):
    return next(t for t in a.schema.tables if t.name == name)
